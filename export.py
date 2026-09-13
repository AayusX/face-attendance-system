import csv
import os

from config import EXPORT_DIR
from nepali_calendar import bs_date_from_ad, today_bs


def _today():
    return today_bs()


def _bs_date(text):
    return bs_date_from_ad(text).isoformat() if bs_date_from_ad(text) is not None else text


def export_attendance(rows, day=None):
    day = day or _today()
    os.makedirs(EXPORT_DIR, exist_ok=True)
    path = EXPORT_DIR / f"attendance_{day}.csv"
    _write_csv(path, rows)
    return path


def export_attendance_range(rows, start, end):
    os.makedirs(EXPORT_DIR, exist_ok=True)
    path = EXPORT_DIR / f"attendance_{start}_to_{end}.csv"
    _write_csv(path, rows)
    return path


def _write_csv(path, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        writer.writerow(["student_id", "name", "date (BS)", "time", "status"])
        writer.writerows(
            [
                (r["student_id"], r["name"], _bs_date(r["date"]), r["time"], r["status"])
                for r in rows
            ]
        )


def export_attendance_excel(rows, start, end):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    os.makedirs(EXPORT_DIR, exist_ok=True)
    path = EXPORT_DIR / f"attendance_{start}_to_{end}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"
    headers = ["student_id", "name", "date (BS)", "time", "status"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DCE6F1")
    for row in rows:
        values = [row[h] for h in ("student_id", "name", "date", "time", "status")]
        values[2] = _bs_date(values[2])
        ws.append(values)
    for col_index in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_index)].width = 16
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    wb.save(path)
    return path